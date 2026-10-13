"""Gera o .docx de um livro a partir dos arquivos-fonte em marcação simples.

Marcação (uma linha por bloco; parágrafos separados por linha em branco):
  !CAPA titulo | volume | subtitulo | publico          -> página de capa
  # Texto      -> título de parte (Capítulo N, Conclusão, Anexo...) em nova página
  ## Texto     -> subtítulo do capítulo
  ### Texto    -> seção
  @ILUS nn | legenda                                     -> espaço para ilustração
  > texto      -> citação/destaque em itálico recuado
  - texto      -> item de lista
  ~ texto      -> parágrafo em itálico
  ^ texto      -> parágrafo centralizado
  ---          -> quebra de página
Dentro do texto: **negrito** e *itálico*.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt

FONTE = "Georgia"


def runs(par, texto, italico=False, tamanho=None):
    for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not parte:
            continue
        b = parte.startswith("**") and parte.endswith("**")
        i = not b and parte.startswith("*") and parte.endswith("*") and len(parte) > 1
        t = parte[2:-2] if b else parte[1:-1] if i else parte
        r = par.add_run(t)
        r.bold = b
        r.italic = i or italico
        r.font.name = FONTE
        if tamanho:
            r.font.size = Pt(tamanho)


def quebra(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def gerar(pasta_fonte, saida):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Cm(2.5))
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    linhas = []
    for arq in sorted(Path(pasta_fonte).glob("*.txt")):
        linhas += arq.read_text(encoding="utf-8").splitlines() + [""]

    blocos, atual = [], []
    for ln in linhas:
        if ln.strip() == "":
            if atual:
                blocos.append(" ".join(atual))
                atual = []
        elif ln.startswith(("- ", "#", "@", "!", "---")) or re.match(r"^[a-e]\) ", ln):
            if atual:
                blocos.append(" ".join(atual))
                atual = []
            blocos.append(ln.rstrip())
        else:
            atual.append(ln.strip())
    if atual:
        blocos.append(" ".join(atual))

    primeiro = True
    for b in blocos:
        if b.startswith("!CAPA"):
            titulo, volume, sub, publico = [x.strip() for x in b[5:].split("|")]
            for _ in range(5):
                doc.add_paragraph()
            for txt, tam, neg, ita in ((titulo, 26, True, False), (volume, 18, True, False),
                                       ("", 11, False, False), (sub, 14, False, True),
                                       ("", 11, False, False), (publico, 12, False, False)):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(txt)
                r.bold, r.italic, r.font.size, r.font.name = neg, ita, Pt(tam), FONTE
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, "**[ ESPAÇO PARA ILUSTRAÇÃO — CAPA ]**", tamanho=12)
            for txt, tam, neg, ita in (("Nilton Dias Sampaio Oliveira Junior", 14, True, False),
                                       ('Do mesmo autor de "Como Vender do Jeito de Deus"', 11, False, True)):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(txt)
                r.bold, r.italic, r.font.size, r.font.name = neg, ita, Pt(tam), FONTE
            primeiro = False
        elif b == "---":
            quebra(doc)
        elif b.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.keep_with_next = True
            runs(p, "**" + b[4:] + "**", tamanho=13)
        elif b.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            runs(p, "**" + b[3:] + "**", tamanho=16)
        elif b.startswith("# "):
            if not primeiro:
                quebra(doc)
            primeiro = False
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            runs(p, "**" + b[2:] + "**", tamanho=18)
        elif b.startswith("@ILUS"):
            num, leg = [x.strip() for x in b[5:].split("|", 1)]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            runs(p, f"**[ ESPAÇO PARA ILUSTRAÇÃO — {num} ]**", tamanho=12)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
            runs(p, f"Fig. {num} — {leg}", italico=True, tamanho=10)
        elif b.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            runs(p, b[2:])
        elif b.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.right_indent = Cm(1)
            runs(p, b[2:], italico=True)
        elif b.startswith("~ "):
            runs(doc.add_paragraph(), b[2:], italico=True)
        elif b.startswith("^ "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, b[2:])
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            runs(p, b)
    doc.save(saida)


if __name__ == "__main__":
    gerar(sys.argv[1], sys.argv[2])
